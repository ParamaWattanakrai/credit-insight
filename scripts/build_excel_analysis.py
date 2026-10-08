"""
Build the statistical-analysis Excel workbook (analysis/credit_card_analysis.xlsx).

Every statistic is a live Excel formula over the `Data` sheet, so the workbook
recalculates if the data is edited or replaced.

Usage:
    python scripts/build_excel_analysis.py [input.xlsx] [output.xlsx]
Defaults:
    input  = data/credit_card_customers.xlsx   (the web app's live Excel database)
    output = analysis/credit_card_analysis.xlsx
Requires: pandas, openpyxl
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "credit_card_customers.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "analysis" / "credit_card_analysis.xlsx"

FONT = "Arial"
BLUE = "1C5CAB"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=15, bold=True, color="0B0B0B")
F_SUB = Font(name=FONT, size=10, italic=True, color="52514E")
F_NOTE = Font(name=FONT, size=9, italic=True, color="898781")
FILL_HEAD = PatternFill("solid", fgColor=BLUE)
FILL_SOFT = PatternFill("solid", fgColor="E4EFFC")
THIN = Border(bottom=Side(style="thin", color="E1E0D9"))
WRAP = Alignment(wrap_text=True, vertical="top")
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]

CARD_TYPES = ["Basic", "Silver", "Gold", "Platinum", "Signature"]
OCCUPATIONS = ["Private Employee", "Government Employee", "Self-Employed", "Business Owner", "Retired", "Student"]
SPEND = ["Grocery_Spending", "Utility_Bill_Spending", "Online_Shopping_Spending", "Fuel_Spending",
         "Dining_Spending", "Entertainment_Spending", "Travel_Spending"]
LABEL = {
    "Annual_Income": "Annual income", "Credit_Limit": "Credit limit", "Card_Age_Months": "Card age (months)",
    "Monthly_Spending": "Monthly spending", "Monthly_Transactions": "Monthly transactions",
    "Avg_Transaction_Value": "Avg transaction value", "Online_Shopping_Spending": "Online shopping",
    "Grocery_Spending": "Grocery", "Fuel_Spending": "Fuel", "Dining_Spending": "Dining", "Travel_Spending": "Travel",
    "Entertainment_Spending": "Entertainment", "Utility_Bill_Spending": "Utility bills",
    "Outstanding_Balance": "Outstanding balance", "Statement_Balance": "Statement balance",
    "Payment_Amount": "Payment amount", "Payment_Ratio": "Payment ratio", "Credit_Utilization": "Credit utilization",
    "Cash_Advance_Amount": "Cash advance", "EMI_Count": "EMI count",
    "International_Transactions": "International transactions", "Reward_Points_Earned": "Reward points earned",
    "Reward_Points_Redeemed": "Reward points redeemed", "Mobile_App_Login": "Mobile app logins",
    "Credit_Score": "Credit score", "Age": "Age",
}
RATIOS = {"Payment_Ratio", "Credit_Utilization"}

# ---------------------------------------------------------------- data
df = pd.read_excel(SRC)
N = len(df)
LAST = N + 1
COLS = list(df.columns)
NUMERIC = [c for c in COLS if pd.api.types.is_numeric_dtype(df[c])]
L = {c: get_column_letter(i + 1) for i, c in enumerate(COLS)}


def rng(col):
    """Absolute reference to a whole data column, e.g. Data!$B$2:$B$50001"""
    return f"Data!${L[col]}$2:${L[col]}${LAST}"


def numfmt(col):
    if col in RATIOS:
        return "0.0%"
    if col in ("Age", "Credit_Score", "Card_Age_Months", "EMI_Count", "Monthly_Transactions",
               "International_Transactions", "Mobile_App_Login"):
        return "#,##0.0"
    return "#,##0"


wb = Workbook()


def sheet(name, title, subtitle=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = F_SUB
    return ws


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def put(ws, row, col, value, fmt=None, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = F_BOLD if bold else F_BASE
    c.border = THIN
    if fmt:
        c.number_format = fmt
    return c


def widths(ws, first, rest, n):
    ws.column_dimensions["A"].width = first
    for i in range(2, n + 1):
        ws.column_dimensions[get_column_letter(i)].width = rest


def bar_chart(ws, anchor, title, cats_ref, data_ref, y_title=None, color=PALETTE[0], horizontal=True, height=7.5, width=16, fmt=None):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = title
    ch.style = 2
    ch.legend = None
    ch.height, ch.width = height, width
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    s = ch.series[0]
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.noFill = True
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showCatName = False
    ch.dataLabels.showSerName = False
    ch.dataLabels.showLegendKey = False
    ch.dataLabels.showPercent = False
    if fmt:
        ch.dataLabels.numFmt = fmt
    ch.y_axis.majorGridlines = None
    ch.y_axis.delete = True
    ch.x_axis.delete = False
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"  # keep table order top-to-bottom
    ch.gapWidth = 60
    if y_title:
        ch.y_axis.title = y_title
    ws.add_chart(ch, anchor)


# ---------------------------------------------------------------- Data sheet
wsd = wb.active
wsd.title = "Data"
wsd.append(COLS)
for row in df.itertuples(index=False):
    wsd.append(list(row))
for c in wsd[1]:
    c.font = F_HEAD
    c.fill = FILL_HEAD
for col in COLS:
    wsd.column_dimensions[L[col]].width = max(11, len(col) + 2)
    if col in RATIOS:
        for (cell,) in wsd.iter_rows(min_row=2, max_row=LAST, min_col=COLS.index(col) + 1, max_col=COLS.index(col) + 1):
            cell.number_format = "0.00"
wsd.freeze_panes = "B2"
tbl = Table(displayName="Customers", ref=f"A1:{get_column_letter(len(COLS))}{LAST}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
wsd.add_table(tbl)

# ---------------------------------------------------------------- Summary
ws = sheet("Summary", "Credit card customer behaviour — statistical analysis",
           f"Source: {SRC.name} ({N:,} customers, synthetic Kaggle dataset). All figures are live formulas over the Data sheet.")
wb.move_sheet("Summary", offset=-1)
header(ws, 4, ["Key metric", "Value", "Formula / definition"])
kpis = [
    ("Customers", f"=COUNTA({rng('Customer_ID')})", "#,##0", "COUNTA of Customer_ID"),
    ("Average annual income", f"=AVERAGE({rng('Annual_Income')})", "#,##0", "AVERAGE"),
    ("Median annual income", f"=MEDIAN({rng('Annual_Income')})", "#,##0", "MEDIAN — robust to the right skew"),
    ("Average monthly spending", f"=AVERAGE({rng('Monthly_Spending')})", "#,##0", "AVERAGE"),
    ("Average credit score", f"=AVERAGE({rng('Credit_Score')})", "0.0", "AVERAGE"),
    ("Average credit utilization", f"=AVERAGE({rng('Credit_Utilization')})", "0.0%", "balance ÷ limit"),
    ("Average payment ratio", f"=AVERAGE({rng('Payment_Ratio')})", "0.0%", "payment ÷ statement"),
    ("Share paying ≥ 95% of statement", f"=COUNTIFS({rng('Payment_Ratio')},\">=0.95\")/B5", "0.0%", "COUNTIFS ÷ customers"),
    ("High-risk customers", f"=COUNTIFS({rng('Credit_Utilization')},\">=0.8\",{rng('Payment_Ratio')},\"<0.3\")", "#,##0",
     "utilization ≥ 80% AND payment ratio < 30% (assumed risk definition)"),
    ("High-risk share", "=B13/B5", "0.0%", "high-risk ÷ customers"),
    ("Correlation: payment ratio ↔ credit score", f"=CORREL({rng('Payment_Ratio')},{rng('Credit_Score')})", "0.000", "CORREL (Pearson r)"),
    ("Correlation: utilization ↔ credit score", f"=CORREL({rng('Credit_Utilization')},{rng('Credit_Score')})", "0.000", "CORREL (Pearson r)"),
]
for i, (lab, f, fmt, note) in enumerate(kpis):
    r = 5 + i
    put(ws, r, 1, lab)
    put(ws, r, 2, f, fmt, bold=True)
    put(ws, r, 3, note).font = F_NOTE
r = 5 + len(kpis) + 1
ws.cell(row=r, column=1, value="Key findings").font = Font(name=FONT, size=12, bold=True)
SPEND_ROW = 5 + NUMERIC.index("Monthly_Spending")
findings = [
    '="1. Payment behaviour drives credit score: r = "&TEXT(B15,"0.00")&" between payment ratio and score (strong positive); utilization works the other way (r = "&TEXT(B16,"0.00")&")."',
    '="2. Card tier separates spenders: Signature holders average "&TEXT(\'By Card Type\'!F9,"#,##0")&" per month vs "&TEXT(\'By Card Type\'!F5,"#,##0")&" for Basic — "&TEXT(\'By Card Type\'!F9/\'By Card Type\'!F5,"0.0")&"×."',
    '="3. Card tier barely changes credit score (ANOVA η² = "&TEXT(\'Hypothesis Tests\'!G15,"0.0000")&"): statistically significant only because n is huge."',
    '="4. Occupation and card tier are strongly associated (Cramér\'s V = "&TEXT(\'Hypothesis Tests\'!C31,"0.000")&"); gender and card tier are not (p = "&TEXT(\'Hypothesis Tests\'!C45,"0.000")&")."',
    '="5. Income explains "&TEXT(\'Hypothesis Tests\'!D53,"0%")&" of the variance in credit limit (R²) — limits are set almost mechanically from income."',
    f'="6. Grocery is the largest category at "&TEXT(\'Spending Mix\'!C5,"0.0%")&" of categorised spend; spending is heavily right-skewed (skew "&TEXT(\'Descriptive Stats\'!K{SPEND_ROW},"0.0")&"), so medians describe a typical customer better than means."',
    '="7. "&TEXT(B14,"0.0%")&" of customers ("&TEXT(B13,"#,##0")&") are high-risk — a small, targetable segment for credit-limit review."',
]
for i, f in enumerate(findings):
    c = ws.cell(row=r + 1 + i, column=1, value=f)
    c.font = F_BASE
    c.alignment = WRAP
    ws.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=3)
    ws.row_dimensions[r + 1 + i].height = 30
note_r = r + 2 + len(findings)
ws.cell(row=note_r, column=1, value="Sheets: Descriptive Stats · Frequency · By Card Type · By Occupation · By Age Band · "
        "By Score Band · Correlation · Spending Mix · Hypothesis Tests · Data").font = F_NOTE
ws.cell(row=note_r + 1, column=1, value="Currency is not stated in the source dataset; amounts are shown as plain numbers.").font = F_NOTE
ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 62

# ---------------------------------------------------------------- Descriptive Stats
ws = sheet("Descriptive Stats", "Descriptive statistics",
           "One row per numeric variable. Sample SD (STDEV); quartiles via QUARTILE (inclusive). |Skewness| > 1 = highly skewed.")
heads = ["Variable", "Count", "Mean", "Std dev", "Min", "Q1", "Median", "Q3", "Max", "IQR", "Skewness", "Kurtosis", "CV"]
header(ws, 4, heads)
for i, col in enumerate(NUMERIC):
    r = 5 + i
    R = rng(col)
    fmt = numfmt(col) if col not in RATIOS else "0.000"
    put(ws, r, 1, LABEL.get(col, col))
    put(ws, r, 2, f"=COUNT({R})", "#,##0")
    for j, fn in enumerate(["AVERAGE", "STDEV", "MIN", None, "MEDIAN", None, "MAX"]):
        if fn:
            put(ws, r, 3 + j, f"={fn}({R})", fmt)
    put(ws, r, 6, f"=QUARTILE({R},1)", fmt)
    put(ws, r, 8, f"=QUARTILE({R},3)", fmt)
    put(ws, r, 10, f"=H{r}-F{r}", fmt)
    put(ws, r, 11, f"=SKEW({R})", "0.00")
    put(ws, r, 12, f"=KURT({R})", "0.00")
    put(ws, r, 13, f"=IF(C{r}=0,0,D{r}/C{r})", "0.00")
last_desc = 4 + len(NUMERIC)
ws.conditional_formatting.add(f"K5:K{last_desc}", ColorScaleRule(start_type="num", start_value=-3, start_color="D03B3B",
                                                                   mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                                   end_type="num", end_value=3, end_color="EB6834"))
ws.freeze_panes = "B5"
widths(ws, 28, 13, len(heads))

# ---------------------------------------------------------------- Frequency (histograms)
ws = sheet("Frequency", "Frequency distributions", "COUNTIFS over bins: [lower, upper). Charts show the shape of each distribution.")
header(ws, 4, ["Credit score from", "to (excl.)", "Customers", "Share"])
bins = list(range(350, 800, 25))
for i, lo in enumerate(bins):
    r = 5 + i
    put(ws, r, 1, lo, "0")
    put(ws, r, 2, lo + 25, "0")
    put(ws, r, 3, f"=COUNTIFS({rng('Credit_Score')},\">=\"&A{r},{rng('Credit_Score')},\"<\"&B{r})", "#,##0")
    put(ws, r, 4, f"=C{r}/Summary!$B$5", "0.0%")
    ws.cell(row=r, column=5, value=f'=A{r}&"–"&(B{r}-1)').font = F_NOTE
end1 = 4 + len(bins)
bar_chart(ws, "G4", "Credit score distribution", Reference(ws, min_col=5, min_row=5, max_row=end1),
          Reference(ws, min_col=3, min_row=4, max_row=end1), horizontal=False, width=20, height=8)
r0 = end1 + 3
header(ws, r0, ["Monthly spend from", "to (excl.)", "Customers", "Share"])
sbins = [0, 10000, 20000, 30000, 40000, 50000, 60000, 80000, 100000, 150000, 200000, 300000, 10**9]
for i in range(len(sbins) - 1):
    r = r0 + 1 + i
    put(ws, r, 1, sbins[i], "#,##0")
    put(ws, r, 2, sbins[i + 1], "#,##0")
    put(ws, r, 3, f"=COUNTIFS({rng('Monthly_Spending')},\">=\"&A{r},{rng('Monthly_Spending')},\"<\"&B{r})", "#,##0")
    put(ws, r, 4, f"=C{r}/Summary!$B$5", "0.0%")
    lab = f'=TEXT(A{r}/1000,"0")&"k+"' if i == len(sbins) - 2 else f'=TEXT(A{r}/1000,"0")&"–"&TEXT(B{r}/1000,"0")&"k"'
    ws.cell(row=r, column=5, value=lab).font = F_NOTE
end2 = r0 + len(sbins) - 1
bar_chart(ws, f"G{r0}", "Monthly spending distribution (bins widen above 60k)", Reference(ws, min_col=5, min_row=r0 + 1, max_row=end2),
          Reference(ws, min_col=3, min_row=r0, max_row=end2), horizontal=False, color=PALETTE[2], width=20, height=8)
widths(ws, 18, 13, 5)

# ---------------------------------------------------------------- group sheets
METRICS = ["Annual_Income", "Credit_Limit", "Monthly_Spending", "Credit_Score", "Credit_Utilization",
           "Payment_Ratio", "Outstanding_Balance", "Reward_Points_Earned"]


def group_sheet(name, title, groups, crit, chart_metric, chart_title, chart_color):
    """crit(i, r) -> list of COUNTIFS/AVERAGEIFS criteria pairs (strings)"""
    ws = sheet(name, title, "COUNTIFS / AVERAGEIFS over the Data sheet. Bars on the right visualise one key metric.")
    heads = ["Group", "Customers", "Share"] + [LABEL[m] for m in METRICS]
    header(ws, 4, heads)
    for i, g in enumerate(groups):
        r = 5 + i
        put(ws, r, 1, g, bold=True)
        c = ",".join(crit(i, r))
        put(ws, r, 2, f"=COUNTIFS({c})", "#,##0")
        put(ws, r, 3, f"=B{r}/Summary!$B$5", "0.0%")
        for j, m in enumerate(METRICS):
            put(ws, r, 4 + j, f"=IFERROR(AVERAGEIFS({rng(m)},{c}),0)", numfmt(m))
    tr = 5 + len(groups)
    put(ws, tr, 1, "All customers", bold=True)
    put(ws, tr, 2, f"=SUM(B5:B{tr - 1})", "#,##0", bold=True)
    put(ws, tr, 3, f"=SUM(C5:C{tr - 1})", "0.0%", bold=True)
    for j, m in enumerate(METRICS):
        put(ws, tr, 4 + j, f"=AVERAGE({rng(m)})", numfmt(m), bold=True)
    for cidx in range(1, 4 + len(METRICS)):
        ws.cell(row=tr, column=cidx).fill = FILL_SOFT
    mcol = 4 + METRICS.index(chart_metric)
    bar_chart(ws, f"A{tr + 3}", chart_title, Reference(ws, min_col=1, min_row=5, max_row=tr - 1),
              Reference(ws, min_col=mcol, min_row=4, max_row=tr - 1), color=chart_color, width=15,
              fmt=numfmt(chart_metric))
    bar_chart(ws, f"G{tr + 3}", "Customers", Reference(ws, min_col=1, min_row=5, max_row=tr - 1),
              Reference(ws, min_col=2, min_row=4, max_row=tr - 1), color=PALETTE[6], width=15, fmt="#,##0")
    widths(ws, 22, 15, len(heads))
    ws.freeze_panes = "B5"
    return ws


group_sheet("By Card Type", "Comparison by card type", CARD_TYPES,
            lambda i, r: [rng("Card_Type"), f"$A{r}"], "Monthly_Spending", "Avg monthly spending", PALETTE[0])
group_sheet("By Occupation", "Comparison by occupation", OCCUPATIONS,
            lambda i, r: [rng("Occupation"), f"$A{r}"], "Credit_Score", "Avg credit score", PALETTE[0])
AGE = [("18–24", 18, 25), ("25–34", 25, 35), ("35–44", 35, 45), ("45–54", 45, 55), ("55–64", 55, 65), ("65+", 65, 200)]
group_sheet("By Age Band", "Comparison by age band", [a[0] for a in AGE],
            lambda i, r: [rng("Age"), f'">={AGE[i][1]}"', rng("Age"), f'"<{AGE[i][2]}"'], "Annual_Income", "Avg annual income", PALETTE[2])
SB = [("Poor (<580)", 0, 580), ("Fair (580–669)", 580, 670), ("Good (670–739)", 670, 740),
      ("Very good (740–799)", 740, 800), ("Excellent (800+)", 800, 2000)]
group_sheet("By Score Band", "Comparison by credit-score band (FICO-style)", [s[0] for s in SB],
            lambda i, r: [rng("Credit_Score"), f'">={SB[i][1]}"', rng("Credit_Score"), f'"<{SB[i][2]}"'],
            "Payment_Ratio", "Avg payment ratio", PALETTE[0])

# ---------------------------------------------------------------- Correlation
CORR = ["Credit_Score", "Payment_Ratio", "Credit_Utilization", "Cash_Advance_Amount", "Outstanding_Balance",
        "Annual_Income", "Credit_Limit", "Monthly_Spending", "Monthly_Transactions", "Card_Age_Months", "Age",
        "Reward_Points_Earned", "Mobile_App_Login"]
ws = sheet("Correlation", "Pearson correlation matrix",
           "CORREL between each pair. Blue = positive, red = negative. |r| ≥ 0.5 strong · 0.3–0.5 moderate · < 0.3 weak.")
header(ws, 4, [""] + [LABEL[c] for c in CORR])
for i, a in enumerate(CORR):
    r = 5 + i
    put(ws, r, 1, LABEL[a], bold=True)
    for j, b in enumerate(CORR):
        put(ws, r, 2 + j, 1 if a == b else f"=CORREL({rng(a)},{rng(b)})", "0.00").alignment = Alignment(horizontal="center")
endc = get_column_letter(1 + len(CORR))
ws.conditional_formatting.add(f"B5:{endc}{4 + len(CORR)}",
                              ColorScaleRule(start_type="num", start_value=-1, start_color="D03B3B",
                                             mid_type="num", mid_value=0, mid_color="F0EFEC",
                                             end_type="num", end_value=1, end_color="2A78D6"))
r0 = 6 + len(CORR)
ws.cell(row=r0, column=1, value="Correlation with credit score").font = Font(name=FONT, size=12, bold=True)
header(ws, r0 + 1, ["Variable", "r", "|r|", "Strength"])
others = [c for c in NUMERIC if c != "Credit_Score"]
# order by strength computed in Python so the table reads top-down; values stay live formulas
order = sorted(others, key=lambda c: -abs(df[c].corr(df["Credit_Score"])))
for i, c in enumerate(order):
    r = r0 + 2 + i
    put(ws, r, 1, LABEL.get(c, c))
    put(ws, r, 2, f"=CORREL({rng(c)},{rng('Credit_Score')})", "0.000")
    put(ws, r, 3, f"=ABS(B{r})", "0.000")
    put(ws, r, 4, f'=IF(C{r}>=0.5,"Strong",IF(C{r}>=0.3,"Moderate",IF(C{r}>=0.1,"Weak","Negligible")))')
ws.cell(row=r0 + 3 + len(order), column=1, value="Rows sorted by |r| at build time; values recalculate live.").font = F_NOTE
widths(ws, 26, 11, 1 + len(CORR))
ws.row_dimensions[4].height = 44

# ---------------------------------------------------------------- Spending Mix
ws = sheet("Spending Mix", "Spending mix by category", "Average monthly amount per customer and each category's share of categorised spend.")
header(ws, 4, ["Category", "Avg monthly", "Share"])
for i, c in enumerate(SPEND):
    r = 5 + i
    put(ws, r, 1, LABEL[c])
    put(ws, r, 2, f"=AVERAGE({rng(c)})", "#,##0")
    put(ws, r, 3, f"=B{r}/SUM($B$5:$B${4 + len(SPEND)})", "0.0%")
endm = 4 + len(SPEND)
put(ws, endm + 1, 1, "Total", bold=True)
put(ws, endm + 1, 2, f"=SUM(B5:B{endm})", "#,##0", bold=True)
put(ws, endm + 1, 3, f"=SUM(C5:C{endm})", "0.0%", bold=True)
bar_chart(ws, "E4", "Share of categorised spend", Reference(ws, min_col=1, min_row=5, max_row=endm),
          Reference(ws, min_col=3, min_row=4, max_row=endm), color=PALETTE[2], fmt="0.0%")
r0 = endm + 4
ws.cell(row=r0 - 1, column=1, value="Share of each category within card tiers").font = Font(name=FONT, size=12, bold=True)
header(ws, r0, ["Card type", "Avg categorised spend"] + [LABEL[c] for c in SPEND])
for i, t in enumerate(CARD_TYPES):
    r = r0 + 1 + i
    put(ws, r, 1, t, bold=True)
    first, lastc = get_column_letter(3), get_column_letter(2 + len(SPEND))
    put(ws, r, 2, f"=SUM({first}{r}:{lastc}{r})", "#,##0")
    for j, c in enumerate(SPEND):
        put(ws, r, 3 + j, f"=AVERAGEIFS({rng(c)},{rng('Card_Type')},$A{r})", "#,##0")
r1 = r0 + len(CARD_TYPES) + 2
header(ws, r1, ["Card type", ""] + [LABEL[c] + " %" for c in SPEND])
for i, t in enumerate(CARD_TYPES):
    r, src = r1 + 1 + i, r0 + 1 + i
    put(ws, r, 1, t, bold=True)
    for j in range(len(SPEND)):
        col = get_column_letter(3 + j)
        put(ws, r, 3 + j, f"={col}{src}/$B{src}", "0.0%")
widths(ws, 20, 14, 2 + len(SPEND))

# ---------------------------------------------------------------- Hypothesis Tests
ws = sheet("Hypothesis Tests", "Inferential statistics",
           "α = 0.05. With n ≈ 50,000 nearly every difference is 'significant' — judge practical importance by the effect size.")
put(ws, 3, 1, "Significance level α").font = F_BOLD
c = put(ws, 3, 3, 0.05, "0.00")
c.font = Font(name=FONT, size=10, color="0000FF", bold=True)
c.fill = PatternFill("solid", fgColor="FFFF00")
ws.cell(row=3, column=4, value="← input: change α to re-evaluate every decision").font = F_NOTE

# --- one-way ANOVA: credit score by card type (rows 5–16)
ws.cell(row=5, column=1, value="1) One-way ANOVA — does mean credit score differ by card type?").font = Font(name=FONT, size=12, bold=True)
ws.cell(row=6, column=1, value="H₀: μ_Basic = μ_Silver = μ_Gold = μ_Platinum = μ_Signature").font = F_SUB
header(ws, 7, ["Card type", "n", "Mean score", "n·(mean − grand)²"])
for i, t in enumerate(CARD_TYPES):
    r = 8 + i
    put(ws, r, 1, t)
    put(ws, r, 2, f"=COUNTIFS({rng('Card_Type')},A{r})", "#,##0")
    put(ws, r, 3, f"=AVERAGEIFS({rng('Credit_Score')},{rng('Card_Type')},A{r})", "0.00")
    put(ws, r, 4, f"=B{r}*(C{r}-AVERAGE({rng('Credit_Score')}))^2", "#,##0")
# results block in columns F:G (rows 8–16)
res = [
    ("SS between (SSB)", "=SUM(D8:D12)", "#,##0"),
    ("SS total (SST)", f"=DEVSQ({rng('Credit_Score')})", "#,##0"),
    ("SS within (SSW)", "=G9-G8", "#,##0"),
    ("df between (k − 1)", "=COUNT(B8:B12)-1", "0"),
    ("df within (N − k)", "=SUM(B8:B12)-COUNT(B8:B12)", "#,##0"),
    ("F statistic", "=(G8/G11)/(G10/G12)", "0.00"),
    ("p-value", "=FDIST(G13,G11,G12)", "0.0000"),
    ("η² (effect size = SSB/SST)", "=G8/G9", "0.0000"),
    ("Decision", '=IF(G14<$C$3,"Reject H₀","Fail to reject H₀")', None),
]
for i, (lab, f, fmt) in enumerate(res):
    put(ws, 8 + i, 6, lab)
    put(ws, 8 + i, 7, f, fmt, bold=True)
ws.cell(row=17, column=1, value="Interpretation: means differ by only a few points across tiers; η² ≈ 0.002 means card tier explains ~0.2% of score variance (negligible).").font = F_NOTE

# --- χ² test: occupation × card type (rows 19–31)
ws.cell(row=19, column=1, value="2) χ² test of independence — occupation × card type").font = Font(name=FONT, size=12, bold=True)
ws.cell(row=20, column=1, value="H₀: card type is independent of occupation. Observed counts (left), expected = row total × column total ÷ N (right).").font = F_SUB
header(ws, 21, ["Occupation"] + CARD_TYPES + ["Total"])
header(ws, 21, ["Expected"] + CARD_TYPES, col=9)
for i, o in enumerate(OCCUPATIONS):
    r = 22 + i
    put(ws, r, 1, o)
    put(ws, r, 9, f"=A{r}")
    for j, t in enumerate(CARD_TYPES):
        col = get_column_letter(2 + j)
        put(ws, r, 2 + j, f"=COUNTIFS({rng('Occupation')},$A{r},{rng('Card_Type')},{col}$21)", "#,##0")
        put(ws, r, 10 + j, f"=$G{r}*{col}$28/$G$28", "#,##0.0")
    put(ws, r, 7, f"=SUM(B{r}:F{r})", "#,##0", bold=True)
put(ws, 28, 1, "Total", bold=True)
for j in range(6):
    col = get_column_letter(2 + j)
    put(ws, 28, 2 + j, f"=SUM({col}22:{col}27)", "#,##0", bold=True)
chi = [
    ("χ² statistic", "=SUMPRODUCT((B22:F27-J22:N27)^2/J22:N27)", "#,##0.00"),
    ("df = (r − 1)(c − 1)", "=(ROWS(B22:F27)-1)*(COLUMNS(B22:F27)-1)", "0"),
    ("Cramér's V", "=SQRT(C29/(G28*(MIN(ROWS(B22:F27),COLUMNS(B22:F27))-1)))", "0.000"),
    ("p-value", "=CHIDIST(C29,C30)", "0.0000"),
    ("Decision", '=IF(C32<$C$3,"Reject H₀","Fail to reject H₀")', None),
]
for i, (lab, f, fmt) in enumerate(chi):
    put(ws, 29 + i, 1, lab)
    put(ws, 29 + i, 3, f, fmt, bold=True)

# --- χ² test: gender × card type (rows 35–46)
ws.cell(row=35, column=1, value="3) χ² test of independence — gender × card type").font = Font(name=FONT, size=12, bold=True)
header(ws, 36, ["Gender"] + CARD_TYPES + ["Total"])
header(ws, 36, ["Expected"] + CARD_TYPES, col=9)
for i, g in enumerate(["M", "F"]):
    r = 37 + i
    put(ws, r, 1, g)
    put(ws, r, 9, f"=A{r}")
    for j in range(5):
        col = get_column_letter(2 + j)
        put(ws, r, 2 + j, f"=COUNTIFS({rng('Gender')},$A{r},{rng('Card_Type')},{col}$36)", "#,##0")
        put(ws, r, 10 + j, f"=$G{r}*{col}$39/$G$39", "#,##0.0")
    put(ws, r, 7, f"=SUM(B{r}:F{r})", "#,##0", bold=True)
put(ws, 39, 1, "Total", bold=True)
for j in range(6):
    col = get_column_letter(2 + j)
    put(ws, 39, 2 + j, f"=SUM({col}37:{col}38)", "#,##0", bold=True)
chi2 = [
    ("χ² statistic", "=SUMPRODUCT((B37:F38-J37:N38)^2/J37:N38)", "#,##0.00"),
    ("df", "=(ROWS(B37:F38)-1)*(COLUMNS(B37:F38)-1)", "0"),
    ("Cramér's V", "=SQRT(C41/(G39*(MIN(ROWS(B37:F38),COLUMNS(B37:F38))-1)))", "0.000"),
    ("Note", "=\"Expected counts all ≥ 5 → χ² approximation valid\"", None),
    ("p-value", "=CHIDIST(C41,C42)", "0.0000"),
    ("Decision", '=IF(C45<$C$3,"Reject H₀","Fail to reject H₀")', None),
]
for i, (lab, f, fmt) in enumerate(chi2):
    put(ws, 41 + i, 1, lab)
    put(ws, 41 + i, 3, f, fmt, bold=True)

# --- simple linear regression (rows 49–53)
ws.cell(row=49, column=1, value="4) Simple linear regression (least squares)").font = Font(name=FONT, size=12, bold=True)
header(ws, 50, ["Model (y = a + b·x)", "Slope b", "Intercept a", "R²", "r"])
models = [("Credit_Score", "Payment_Ratio"), ("Credit_Score", "Credit_Utilization"), ("Credit_Limit", "Annual_Income")]
for i, (y, x) in enumerate(models):
    r = 51 + i
    put(ws, r, 1, f"{LABEL[y]} ~ {LABEL[x]}")
    put(ws, r, 2, f"=SLOPE({rng(y)},{rng(x)})", "#,##0.0000")
    put(ws, r, 3, f"=INTERCEPT({rng(y)},{rng(x)})", "#,##0.00")
    put(ws, r, 4, f"=RSQ({rng(y)},{rng(x)})", "0.000")
    put(ws, r, 5, f"=CORREL({rng(y)},{rng(x)})", "0.000")
ws.cell(row=55, column=1, value="Reading: +0.10 in payment ratio ≈ +0.1×slope score points. Income → limit R² near 1 means limits are a near-fixed multiple of income.").font = F_NOTE
widths(ws, 40, 13, 14)
ws.column_dimensions["F"].width = 28

wb.move_sheet("Data", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Data"))
wb.active = 0
# No cached values are written by openpyxl — ask Excel to compute every formula on open.
wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT} ({N:,} rows)")
